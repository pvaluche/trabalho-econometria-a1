"""
src/entrega_2t.py
Gera a planilha XLSX oficial de previsoes para o 2o Turno do Desafio FGV EPGE 2026.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from src.arredondamento import maiores_restos
from src.config import OUTPUTS_DIR

CANDIDATOS_2T_EDITAL = [
    "Flávio Bolsonaro",
    "Luiz Inácio Lula da Silva",
]

PARTIDOS_2T_EDITAL = {
    "Flávio Bolsonaro": "PL",
    "Luiz Inácio Lula da Silva": "PT",
}


def _obter_valor_adicional(
    adicionais: dict[str, float],
    chaves: tuple[str, ...],
    nome_campo: str,
) -> float:
    for ch in chaves:
        if ch in adicionais and adicionais[ch] is not None:
            return float(adicionais[ch])
    raise ValueError(
        f"Valor obrigatorio ausente em `adicionais` para '{nome_campo}'. Chaves aceitas: {chaves}."
    )


def gerar_planilha_entrega_2t(
    previsoes_candidatos: dict[str, float],
    adicionais: dict[str, float],
    caminho_saida: Path | str | None = None,
) -> Path:
    """
    Gera a planilha XLSX de entrega do 2o Turno com 2 abas perfeitamente formatadas.

    Aba 1 (`Candidatos`): Flávio Bolsonaro (PL), Luiz Inácio Lula da Silva (PT) e Total (100,0%).
    Aba 2 (`Adicionais`): Abstenção, Votos brancos, Votos nulos.
    Levanta ValueError se faltar qualquer candidato ou variavel adicional.
    """
    if caminho_saida is None:
        caminho_saida = OUTPUTS_DIR / "previsao_2026_2t.xlsx"
    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    valores_ordenados = []
    for c in CANDIDATOS_2T_EDITAL:
        if c not in previsoes_candidatos or previsoes_candidatos[c] is None:
            raise ValueError(f"Candidato obrigatorio ausente em `previsoes_candidatos`: '{c}'")
        valores_ordenados.append(float(previsoes_candidatos[c]))

    val_abst = _obter_valor_adicional(adicionais, ("abstencao", "abstenção"), "Abstenção")
    val_brancos = _obter_valor_adicional(adicionais, ("brancos", "votos brancos"), "Votos brancos")
    val_nulos = _obter_valor_adicional(adicionais, ("nulos", "votos nulos"), "Votos nulos")

    valores_fechados = maiores_restos(valores_ordenados, total=100.0, casas=1)

    wb = openpyxl.Workbook()

    # Aba 1: Candidatos
    ws1 = wb.active
    ws1.title = "Candidatos"
    ws1.append(["Candidato(a)", "Partido", "Previsão de votos válidos (%)"])

    for i, (nome, pct) in enumerate(zip(CANDIDATOS_2T_EDITAL, valores_fechados), start=2):
        ws1.append([nome, PARTIDOS_2T_EDITAL[nome], round(float(pct), 1)])
        ws1.cell(row=i, column=3).number_format = "0.0"

    linha_total_idx = len(CANDIDATOS_2T_EDITAL) + 2
    ws1.append(["Total", "", 100.0])
    ws1.cell(row=linha_total_idx, column=3).number_format = "0.0"

    # Aba 2: Adicionais
    ws2 = wb.create_sheet("Adicionais")
    ws2.append(["Resultado", "Previsão (%)"])

    ws2.append(["Abstenção", round(float(val_abst), 1)])
    ws2.cell(row=2, column=2).number_format = "0.0"

    ws2.append(["Votos brancos", round(float(val_brancos), 1)])
    ws2.cell(row=3, column=2).number_format = "0.0"

    ws2.append(["Votos nulos", round(float(val_nulos), 1)])
    ws2.cell(row=4, column=2).number_format = "0.0"

    wb.save(caminho_saida)
    return caminho_saida
