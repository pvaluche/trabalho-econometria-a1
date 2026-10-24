"""
tests/test_pipeline_2t.py
Suite de testes automatizados para a pipeline do 2o Turno (Desafio FGV EPGE 2026).
"""

from __future__ import annotations

import json

import openpyxl
import pandas as pd
import pytest

from src.backtest_2t import (
    DATAS_CORTE_2T,
    aplicar_ajuste_persistencia_1t,
    calcular_prior_estrutural_1t,
    calcular_taus_historicos,
    carregar_pesquisas_2t_historicas,
    carregar_pesquisas_2026_2t,
    estimar_m0_2t,
    estimar_m1_2t,
    estimar_m2_2t_peso_1t,
)
from src.config import MANUAL_DIR, PROCESSED_DIR, REPORTS_DIR
from src.entrega_2t import gerar_planilha_entrega_2t


def test_arquivos_processados_2t_existem():
    assert (PROCESSED_DIR / "pesquisas_historicas_2t.parquet").exists()
    assert (PROCESSED_DIR / "tse_historico_2t.parquet").exists()
    assert (PROCESSED_DIR / "tse_uf_historico_2t.parquet").exists()
    assert (PROCESSED_DIR / "tse_2026_1t_oficial.json").exists()
    assert (MANUAL_DIR / "pesquisas_2026_2t.csv").exists()
    assert (MANUAL_DIR / "vespera_extra_2026.csv").exists()


def test_vespera_extra_2026_csv_integridade():
    df_extra = pd.read_csv(MANUAL_DIR / "vespera_extra_2026.csv")
    assert set(df_extra["instituto"]) == {"Quaest", "AtlasIntel", "CNT/MDA"}
    for col in ["registro_tse", "fonte_url", "Flávio Bolsonaro", "Luiz Inácio Lula da Silva"]:
        assert col in df_extra.columns
        assert df_extra[col].notna().all()
    assert df_extra["registro_tse"].str.startswith("BR-").all()
    assert df_extra["fonte_url"].str.startswith("https://").all()


def test_zero_vazamento_temporal_pesquisas_2t():
    df_hist = carregar_pesquisas_2t_historicas()
    assert len(df_hist) >= 100
    for ano, grp in df_hist.groupby("eleicao"):
        corte = DATAS_CORTE_2T[int(ano)].isoformat()
        assert (grp["data_divulgacao"] <= corte).all(), f"Vazamento temporal em {ano}"
        assert ((grp["pct_pt_validos"] + grp["pct_adv_validos"] - 100.0).abs() < 1e-6).all()

    df_2026 = carregar_pesquisas_2026_2t()
    assert len(df_2026) >= 20
    assert (df_2026["data_divulgacao"] >= "2026-10-05").all()
    assert (df_2026["data_divulgacao"] <= "2026-10-24").all()
    assert ((df_2026["pct_pt_validos"] + df_2026["pct_adv_validos"] - 100.0).abs() < 1e-6).all()


def test_sanidade_tse_2026_1t_oficial():
    with open(PROCESSED_DIR / "tse_2026_1t_oficial.json", encoding="utf-8") as f:
        d = json.load(f)

    assert abs(d["abstencao_pct"] - 21.0836) < 0.01
    assert abs(d["brancos_pct"] - 1.8366) < 0.01
    assert abs(d["nulos_pct"] - 2.9329) < 0.01

    cands = d["candidatos"]
    assert abs(cands["Flávio Bolsonaro"]["pct_validos"] - 47.0278) < 0.01
    assert abs(cands["Luiz Inácio Lula da Silva"]["pct_validos"] - 45.1628) < 0.01
    soma_validos = sum(v["pct_validos"] for v in cands.values())
    assert abs(soma_validos - 100.0) < 1e-4


def test_modelos_2t_e_taus_historicos():
    df_2026 = carregar_pesquisas_2026_2t()
    p0 = estimar_m0_2t(df_2026)
    p1 = estimar_m1_2t(df_2026, 2026, meia_vida=5.0)
    p2 = estimar_m2_2t_peso_1t(df_2026, 2026, sigma0=1.5)
    p_adj = aplicar_ajuste_persistencia_1t(
        df_2026, 2026, rho=0.20, modo_base="M2_2T", param_base=1.5
    )

    for p in [p0, p1, p2, p_adj]:
        assert 40.0 <= p <= 60.0

    taus = calcular_taus_historicos([2006, 2010, 2014, 2018, 2022])
    assert abs(taus[2006] - 1.2530) < 0.01
    p_prior_med = calcular_prior_estrutural_1t(
        2026, [2006, 2010, 2014, 2018, 2022], estatistica="media"
    )
    p_prior_mdn = calcular_prior_estrutural_1t(
        2026, [2006, 2010, 2014, 2018, 2022], estatistica="mediana"
    )
    assert p_prior_med > p_prior_mdn


def test_planilha_entrega_2t_formato_e_erros_sem_defaults(tmp_path):
    caminho = tmp_path / "teste_2t.xlsx"
    gerar_planilha_entrega_2t(
        {"Flávio Bolsonaro": 51.44, "Luiz Inácio Lula da Silva": 48.56},
        {"abstencao": 21.1, "brancos": 1.8, "nulos": 2.9},
        caminho_saida=caminho,
    )
    assert caminho.exists()

    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Candidatos", "Adicionais"]

    ws1 = wb["Candidatos"]
    assert [ws1.cell(1, c).value for c in range(1, 4)] == [
        "Candidato(a)",
        "Partido",
        "Previsão de votos válidos (%)",
    ]
    assert ws1.cell(2, 1).value == "Flávio Bolsonaro"
    assert ws1.cell(2, 2).value == "PL"
    assert ws1.cell(3, 1).value == "Luiz Inácio Lula da Silva"
    assert ws1.cell(3, 2).value == "PT"
    assert ws1.cell(4, 1).value == "Total"
    assert abs(float(ws1.cell(2, 3).value) + float(ws1.cell(3, 3).value) - 100.0) < 1e-6
    assert float(ws1.cell(4, 3).value) == 100.0

    ws2 = wb["Adicionais"]
    assert [ws2.cell(1, c).value for c in range(1, 3)] == ["Resultado", "Previsão (%)"]
    assert [ws2.cell(r, 1).value for r in range(2, 5)] == [
        "Abstenção",
        "Votos brancos",
        "Votos nulos",
    ]
    for r in range(2, 5):
        val = float(ws2.cell(r, 2).value)
        assert 0.0 <= val <= 100.0

    # Verifica que falta de qualquer campo levanta ValueError (sem defaults hardcoded)
    with pytest.raises(ValueError):
        gerar_planilha_entrega_2t(
            {"Flávio Bolsonaro": 51.3, "Luiz Inácio Lula da Silva": 48.7},
            {"brancos": 1.8, "nulos": 2.9},
            caminho_saida=tmp_path / "err1.xlsx",
        )
    with pytest.raises(ValueError):
        gerar_planilha_entrega_2t(
            {"Flávio Bolsonaro": 51.3},
            {"abstencao": 21.1, "brancos": 1.8, "nulos": 2.9},
            caminho_saida=tmp_path / "err2.xlsx",
        )


def test_relatorio_backtest_2t_gerado_pelo_script():
    rel = (REPORTS_DIR / "relatorio_backtest_2t.md").read_text(encoding="utf-8")
    assert "Arquivo gerado 100% automaticamente por `scripts/executar_backtest_2t.py`" in rel
    assert "Prior transferência com tau médio" in rel
    assert "Prior transferência com tau mediano" in rel
    assert "1,2530" in rel
    assert "7,92%" in rel
    assert "Gerp no peso 1T" not in rel
