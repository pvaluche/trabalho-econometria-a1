"""
scripts/gerar_numeros_latex.py
Le a previsao final do pipeline (outputs/previsao_2026.xlsx e pesquisas_2026.csv)
e reescreve docs/numeros_finais.tex com as definicoes de macros LaTeX (\newcommand).
Zero travessoes em todo o arquivo.
"""

from __future__ import annotations

import argparse
from pathlib import Path
import sys
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.arredondamento import maiores_restos
from src.backtest import estimar_m0
from src.config import CANDIDATOS_EDITAL, DOCS_DIR, MANUAL_DIR, OUTPUTS_DIR

MAPA_MACROS_OFICIAL = {
    "Luiz Inácio Lula da Silva": "pLula",
    "Flávio Bolsonaro": "pFlavio",
    "Augusto Cury": "pCury",
    "Ronaldo Caiado": "pCaiado",
    "Renan Santos": "pRenan",
    "Romeu Zema": "pZema",
    "Samara Martins": "pSamara",
    "Clariana Barão": "pClariana",
    "Edmilson Costa": "pEdmilson",
    "Hertz Dias": "pHertz",
    "Rui Costa Pimenta": "pRui",
    "Wilson Grassi": "pWilson",
}

MAPA_MACROS_M0 = {
    "Luiz Inácio Lula da Silva": "mLula",
    "Flávio Bolsonaro": "mFlavio",
    "Augusto Cury": "mCury",
    "Ronaldo Caiado": "mCaiado",
    "Renan Santos": "mRenan",
    "Romeu Zema": "mZema",
}


def fmt_virgula(val: float, casas: int = 1) -> str:
    """Formata valor numerico com virgula decimal e numero fixo de casas."""
    s = f"{val:.{casas}f}"
    return s.replace(".", ",")


def formatar_lista_pesquisas(df: pd.DataFrame) -> str:
    """Gera string de institutos e protocolos de registro formatada em portugues."""
    grupos: dict[str, list[str]] = {}
    for _, row in df.iterrows():
        inst_raw = str(row["instituto"]).strip()
        if "PoderData" in inst_raw:
            inst = "PoderData"
        elif "Real Time" in inst_raw:
            inst = "Real Time Big Data"
        elif "Datafolha" in inst_raw:
            inst = "Datafolha"
        elif "Quaest" in inst_raw:
            inst = "Quaest"
        elif "Atlas" in inst_raw:
            inst = "AtlasIntel"
        else:
            inst = inst_raw

        reg = str(row["registro_tse"]).strip() if pd.notna(row["registro_tse"]) else ""
        if inst not in grupos:
            grupos[inst] = []
        if reg and reg not in grupos[inst]:
            grupos[inst].append(reg)

    partes = []
    for inst, regs in grupos.items():
        if not regs:
            partes.append(inst)
        elif len(regs) == 1:
            partes.append(f"{inst} ({regs[0]})")
        elif len(regs) == 2:
            partes.append(f"{inst} ({regs[0]} e {regs[1]})")
        else:
            regs_str = ", ".join(regs[:-1]) + f" e {regs[-1]}"
            partes.append(f"{inst} ({regs_str})")

    if not partes:
        return "Nenhuma pesquisa registrada"
    if len(partes) == 1:
        return partes[0]
    return ", ".join(partes[:-1]) + f" e {partes[-1]}"


def carregar_valores_xlsx() -> tuple[dict[str, float], dict[str, float]]:
    """Carrega os valores da Aba 1 e Aba 2 do arquivo outputs/previsao_2026.xlsx."""
    xlsx_path = OUTPUTS_DIR / "previsao_2026.xlsx"
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Arquivo de entrega nao encontrado em {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws1 = wb["Candidatos"]
    ws2 = wb["Adicionais"]

    candidatos_vals = {}
    for row in ws1.iter_rows(min_row=2, values_only=True):
        nome = str(row[0] or "").strip()
        if nome and nome.lower() != "total":
            val = float(row[2])
            candidatos_vals[nome] = val

    aba2_vals = {}
    for row in ws2.iter_rows(min_row=2, values_only=True):
        item = str(row[0] or "").strip()
        if item:
            val = float(row[1])
            aba2_vals[item.lower()] = val

    return candidatos_vals, aba2_vals


def calcular_m0_dict(df_2026: pd.DataFrame) -> dict[str, float]:
    """Calcula os valores do M0 puro fechados a 100,0% por maiores restos."""
    p_m0 = estimar_m0(df_2026, 2026)
    r_m0 = maiores_restos([p_m0[c] for c in CANDIDATOS_EDITAL], total=100.0, casas=1)
    return {c: round(v, 1) for c, v in zip(CANDIDATOS_EDITAL, r_m0)}


def gerar_conteudo_latex(
    status_previsao: str = "FINAL (corte 03/10/2026 20h00)",
) -> str:
    """Gera o texto completo para docs/numeros_finais.tex."""
    csv_path = MANUAL_DIR / "pesquisas_2026.csv"
    df_2026 = pd.read_csv(csv_path)

    n_pesquisas = len(df_2026)
    lista_pesquisas = formatar_lista_pesquisas(df_2026)

    candidatos_vals, aba2_vals = carregar_valores_xlsx()
    m0_dict = calcular_m0_dict(df_2026)

    linhas = [
        "% =====================================================================",
        "% numeros_finais.tex",
        "% Gerado pelo pipeline (scripts/gerar_numeros_latex.py) apos o corte.",
        "% NAO editar a mao: rodar o script e recompilar o PDF.",
        "% =====================================================================",
        f"\\newcommand{{\\statusprevisao}}{{{status_previsao}}}",
        f"\\newcommand{{\\npesquisas}}{{{n_pesquisas}}}",
        f"\\newcommand{{\\listapesquisas}}{{{lista_pesquisas}}}",
        "",
        "% Modelo oficial (votos validos, %)",
    ]

    # Votos validos oficiais (Aba 1)
    for cand, macro in MAPA_MACROS_OFICIAL.items():
        val = candidatos_vals.get(cand, 0.0)
        linhas.append(f"\\newcommand{{\\{macro}}}{{{fmt_virgula(val, 1)}}}")

    linhas.append("")
    linhas.append("% M0 puro (sensibilidade)")
    for cand, macro in MAPA_MACROS_M0.items():
        val = m0_dict.get(cand, 0.0)
        linhas.append(f"\\newcommand{{\\{macro}}}{{{fmt_virgula(val, 1)}}}")

    # Aba 2: Abstencao, Brancos, Nulos
    # Localiza com flexibilidade nas chaves de aba2_vals
    p_abst = 22.3
    p_brancos = 1.6
    p_nulos = 2.8

    for k, v in aba2_vals.items():
        if "abst" in k:
            p_abst = v
        elif "branco" in k:
            p_brancos = v
        elif "nulo" in k:
            p_nulos = v

    linhas.append("")
    linhas.append("% Aba 2")
    linhas.append(f"\\newcommand{{\\pAbst}}{{{fmt_virgula(p_abst, 1)}}}")
    linhas.append(f"\\newcommand{{\\pBrancos}}{{{fmt_virgula(p_brancos, 1)}}}")
    linhas.append(f"\\newcommand{{\\pNulos}}{{{fmt_virgula(p_nulos, 1)}}}")
    linhas.append("")

    return "\n".join(linhas)


def main():
    parser = argparse.ArgumentParser(description="Gera docs/numeros_finais.tex a partir do pipeline.")
    parser.add_argument(
        "--status",
        default="FINAL (corte 03/10/2026 20h00)",
        help="Texto do macro statusprevisao.",
    )
    args = parser.parse_args()

    conteudo = gerar_conteudo_latex(status_previsao=args.status)
    dest = DOCS_DIR / "numeros_finais.tex"
    dest.write_text(conteudo, encoding="utf-8")
    print(f"OK: {dest} gerado com sucesso ({dest.stat().st_size} bytes).")


if __name__ == "__main__":
    main()
