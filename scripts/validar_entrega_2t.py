"""
scripts/validar_entrega_2t.py
Validador estrito da planilha oficial de entrega do 2o Turno (outputs/previsao_2026_2t.xlsx).
"""

from __future__ import annotations

from pathlib import Path
import sys
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.config import OUTPUTS_DIR
from src.entrega_2t import CANDIDATOS_2T_EDITAL, PARTIDOS_2T_EDITAL


def validar_xlsx_2t(caminho: Path | str | None = None) -> list[str]:
    if caminho is None:
        caminho = OUTPUTS_DIR / "previsao_2026_2t.xlsx"
    caminho = Path(caminho)
    erros: list[str] = []

    if not caminho.exists():
        return [f"Arquivo inexistente: {caminho}"]

    wb = openpyxl.load_workbook(caminho, data_only=True)
    if wb.sheetnames != ["Candidatos", "Adicionais"]:
        erros.append(f"Abas invalidas: esperado ['Candidatos', 'Adicionais'], encontrado {wb.sheetnames}")
        return erros

    ws1 = wb["Candidatos"]
    header1 = [ws1.cell(1, c).value for c in range(1, 4)]
    if header1 != ["Candidato(a)", "Partido", "Previsão de votos válidos (%)"]:
        erros.append(f"Cabecalho da Aba 1 invalido: {header1}")

    soma_decimos = 0
    for idx, cand in enumerate(CANDIDATOS_2T_EDITAL, start=2):
        nome = ws1.cell(idx, 1).value
        partido = ws1.cell(idx, 2).value
        val = ws1.cell(idx, 3).value
        fmt = ws1.cell(idx, 3).number_format
        if nome != cand:
            erros.append(f"Linha {idx} Aba 1: esperado candidato '{cand}', encontrado '{nome}'")
        if partido != PARTIDOS_2T_EDITAL[cand]:
            erros.append(f"Linha {idx} Aba 1: esperado partido '{PARTIDOS_2T_EDITAL[cand]}', encontrado '{partido}'")
        if val is None or not (0.0 <= float(val) <= 100.0):
            erros.append(f"Linha {idx} Aba 1: percentual invalido '{val}'")
        else:
            soma_decimos += int(round(float(val) * 10))
        if fmt != "0.0":
            erros.append(f"Linha {idx} Aba 1: formato numerico esperado '0.0', encontrado '{fmt}'")

    if soma_decimos != 1000:
        erros.append(f"Soma em decimos dos candidatos diferente de 1000: {soma_decimos}")

    if ws1.cell(4, 1).value != "Total" or float(ws1.cell(4, 3).value or 0.0) != 100.0:
        erros.append("Linha Total da Aba 1 invalida")

    ws2 = wb["Adicionais"]
    header2 = [ws2.cell(1, c).value for c in range(1, 3)]
    if header2 != ["Resultado", "Previsão (%)"]:
        erros.append(f"Cabecalho da Aba 2 invalido: {header2}")

    esperados_aba2 = ["Abstenção", "Votos brancos", "Votos nulos"]
    for idx, rot in enumerate(esperados_aba2, start=2):
        lbl = ws2.cell(idx, 1).value
        val = ws2.cell(idx, 2).value
        fmt = ws2.cell(idx, 2).number_format
        if lbl != rot:
            erros.append(f"Linha {idx} Aba 2: esperado '{rot}', encontrado '{lbl}'")
        if val is None or not (0.0 <= float(val) <= 100.0):
            erros.append(f"Linha {idx} Aba 2: valor invalido '{val}'")
        if fmt != "0.0":
            erros.append(f"Linha {idx} Aba 2: formato numerico esperado '0.0', encontrado '{fmt}'")

    return erros


if __name__ == "__main__":
    alvo = sys.argv[1] if len(sys.argv) > 1 else None
    erros_encontrados = validar_xlsx_2t(alvo)
    if erros_encontrados:
        for e in erros_encontrados:
            print(f"ERRO: {e}")
        sys.exit(1)
    print("OK: arquivo previsao_2026_2t.xlsx valido para entrega.")
